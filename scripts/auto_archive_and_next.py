#!/usr/bin/env python3
"""独立于 Claude 会话的自动收尾脚本（由系统 crontab 触发）。

职责：当指定补跑任务结束后，自动执行：
  1. 归档成功案例到 archives/successful_cve_0917/
  2. 更新 工作流迪普IPS结果统计0917.xlsx 的 J/K 列 + 统计汇总
  3. 用最优配置启动下一轮补跑（88 个协议未支持的 CVE）

幂等：用 done 标记文件防止重复执行。
"""
import json
import os
import shutil
import subprocess
import time
import glob
import urllib.request

API = "http://127.0.0.1:8000"
ROOT = "/home/ustc/cve-hunter"
DONE_MARK = f"{ROOT}/output/auto_done_5a3f032683ca.txt"
NEXT_LIST = "/tmp/protocol_unsupported.txt"


def log(msg):
    line = f"{time.strftime('%Y-%m-%d %H:%M:%S')} {msg}"
    with open(f"{ROOT}/output/auto_archive.log", "a") as f:
        f.write(line + "\n")
    print(line)


def task_status(task_id):
    try:
        with urllib.request.urlopen(f"{API}/api/tasks/{task_id}", timeout=10) as r:
            d = json.load(r)
        return d.get("status")
    except Exception as e:
        return None


def archive_success():
    latest = {}
    for rj in glob.glob(f"{ROOT}/output/web_runs/*/*/result.json"):
        cve = os.path.basename(os.path.dirname(rj)).upper()
        if not cve.startswith("CVE-"):
            continue
        try:
            d = json.load(open(rj))
        except Exception:
            continue
        ts = d.get("timestamp", "")
        if cve not in latest or ts > latest[cve][0]:
            latest[cve] = (ts, os.path.dirname(rj), d.get("status", ""))
    succ = {c: v for c, v in latest.items() if v[2] == "SUCCESS"}
    arch = f"{ROOT}/archives/successful_cve_0917"
    os.makedirs(arch, exist_ok=True)
    for cve, (ts, d2, st) in sorted(succ.items()):
        dst = os.path.join(arch, cve)
        os.makedirs(dst, exist_ok=True)
        shutil.copy2(os.path.join(d2, "result.json"), os.path.join(dst, "result.json"))
        for fn in os.listdir(d2):
            if fn != "result.json" and (fn.endswith(".pcap") or fn == "poc.http"):
                shutil.copy2(os.path.join(d2, fn), os.path.join(dst, fn))
    with open(os.path.join(arch, "README.md"), "w") as f:
        f.write(f"# 成功 CVE 归档（自动）\n\n共 {len(succ)} 个成功 CVE。\n\n")
        for c in sorted(succ):
            f.write(f"- {c}\n")
    return latest, len(succ)


def update_xlsx(latest):
    import openpyxl
    wb = openpyxl.load_workbook(f"{ROOT}/工作流迪普IPS结果统计0917.xlsx")
    ws = wb["迪普IPS待补充CVE运行结果"]
    h3c_n = dptech_n = both_n = any_n = 0
    for r in range(2, ws.max_row + 1):
        cve = ws.cell(row=r, column=3).value
        if not cve:
            continue
        cve = str(cve).strip().upper()
        v = latest.get(cve)
        if not v:
            continue
        d2 = v[1]
        h3c = dptech = False
        try:
            dd = json.load(open(os.path.join(d2, "result.json")))
            for m in dd.get("cve_ips_match_details", []):
                src = str(m.get("source", "")).lower()
                if src == "h3c":
                    h3c = True
                elif src == "dptech":
                    dptech = True
        except Exception:
            pass
        ws.cell(row=r, column=10, value=("是" if h3c else ""))
        ws.cell(row=r, column=11, value=("是" if dptech else ""))
        if h3c:
            h3c_n += 1
        if dptech:
            dptech_n += 1
        if h3c and dptech:
            both_n += 1
        if h3c or dptech:
            any_n += 1
    st = wb["统计"]
    st.cell(row=2, column=2, value=any_n)
    st.cell(row=3, column=2, value=h3c_n)
    st.cell(row=4, column=2, value=dptech_n)
    st.cell(row=5, column=2, value=both_n)
    wb.save(f"{ROOT}/工作流迪普IPS结果统计0917.xlsx")
    return any_n, h3c_n, dptech_n, both_n


def start_next_round():
    if not os.path.isfile(NEXT_LIST):
        return None
    # 避免重复启动：检查是否有 running/queued 的任务
    try:
        with urllib.request.urlopen(f"{API}/api/tasks", timeout=10) as r:
            tasks = json.load(r)
    except Exception:
        return None
    for t in tasks:
        if t.get("status") in ("queued", "running", "stopping") and t.get("id") != "5a3f032683ca":
            return None  # 已有别的任务在跑，不重复启动
    # 启动下一轮
    import subprocess
    boundary = "----AutoBoundary"
    body = []
    with open(NEXT_LIST) as f:
        content = f.read()
    fields = {
        "mode": "full", "concurrency": "12", "target_ip": "10.55.1.21",
        "local_only": "false", "allow_reference_links": "true",
        "docker_enabled": "false", "environment_discovery": "false",
    }
    for name, value in fields.items():
        body.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n')
    body.append(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="next.txt"\r\nContent-Type: text/plain\r\n\r\n{content}\r\n')
    body.append(f'--{boundary}--\r\n')
    data = "".join(body).encode("utf-8")
    req = urllib.request.Request(
        f"{API}/api/tasks",
        data=data,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            d = json.load(r)
        return d.get("id")
    except Exception as e:
        log(f"启动下一轮失败: {e}")
        return None


def main():
    if os.path.isfile(DONE_MARK):
        return  # 已执行过
    status = task_status("5a3f032683ca")
    if status in (None, "queued", "running", "stopping"):
        return  # 还在跑
    log(f"任务 5a3f032683ca 已结束(status={status})，开始归档+xlsx更新")
    latest, n_succ = archive_success()
    any_n, h3c_n, dptech_n, both_n = update_xlsx(latest)
    log(f"归档成功 {n_succ} 个；xlsx 总命中 {any_n}/h3c {h3c_n}/迪普 {dptech_n}/重叠 {both_n}")
    nid = start_next_round()
    if nid:
        log(f"已启动下一轮任务 {nid}")
    with open(DONE_MARK, "w") as f:
        f.write(f"done at {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
    log("收尾完成")


if __name__ == "__main__":
    main()
